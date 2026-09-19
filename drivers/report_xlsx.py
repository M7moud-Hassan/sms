"""
Shared Excel (openpyxl) helpers that mirror the ReportLab PDF reports:
same brand-coloured header row, zebra striping, thin grey grid, RTL sheet
and a totals block.
 
Place this file where every reports app can import it (e.g. a shared
`core` package, or a copy inside each app) and import as:  from . import report_xlsx as rx
 
NOTE on Arabic: unlike the PDF builders, do NOT reshape/bidi the text here.
Excel shapes Arabic itself, and `sheet_view.rightToLeft = True` flips the
column order, so we write each row REVERSED (the PDF's rightmost/first-read
column becomes Excel column A).
 
Dependency:  pip install openpyxl
"""
import io
import zipfile
 
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
 
FONT = "Amiri"          # falls back to a system Arabic font if not installed
GRID = "D0D0D0"
ZEBRA = "F4F7FB"
WHITE = "FFFFFF"
BLACK = "000000"
 
PDF_CT = "application/pdf"
XLSX_CT = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
 
 
def _hex(c):
    return str(c).lstrip("#").upper()
 
 
def _border():
    s = Side(style="thin", color=GRID)
    return Border(left=s, right=s, top=s, bottom=s)
 
 
def new_sheet(title="report"):
    wb = Workbook()
    ws = wb.active
    ws.title = (title or "report")[:31]
    ws.sheet_view.rightToLeft = True
    return wb, ws
 
 
def fmt_dt(dt):
    return dt.strftime("%Y-%m-%d %H:%M") if hasattr(dt, "strftime") else ("" if dt is None else str(dt))
 
 
def fmt_date(d):
    return d.strftime("%Y-%m-%d") if hasattr(d, "strftime") else ("" if d is None else str(d))
 
 
def write_title(ws, brand, title, date_from, date_to, ncols, start_row=1):
    brand = _hex(brand)
    ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row, end_column=ncols)
    t = ws.cell(row=start_row, column=1, value=title)
    t.font = Font(name=FONT, bold=True, size=16, color=brand)
    t.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[start_row].height = 26
 
    ws.merge_cells(start_row=start_row + 1, start_column=1, end_row=start_row + 1, end_column=ncols)
    d = ws.cell(row=start_row + 1, column=1, value=f"من {date_from} إلى {date_to}")
    d.font = Font(name=FONT, size=11, color="555555")
    d.alignment = Alignment(horizontal="right", vertical="center")
    return start_row + 3          # leave one blank row before the table
 
 
def write_table(ws, brand, header, rows, col_widths=None, start_row=1,
                number_cols=None, zebra=True, total_row=False, font_size=12):
    """
    `header`/`rows` are given in PDF (left->right) order and written REVERSED
    so column A ends up on the right (RTL), matching the PDF layout.
    `number_cols` = set of PDF-order indices to format as #,##0.00.
    `total_row`   = style the LAST data row as a brand-coloured totals row.
    Returns the next free row.
    """
    brand = _hex(brand)
    ncols = len(header)
    border = _border()
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    number_cols = number_cols or set()
 
    # header
    for j, val in enumerate(reversed(header), start=1):
        c = ws.cell(row=start_row, column=j, value=val)
        c.font = Font(name=FONT, bold=True, color=WHITE, size=font_size)
        c.fill = PatternFill(fill_type="solid", fgColor=brand)
        c.alignment = center
        c.border = border
 
    for i, row in enumerate(rows):
        r = start_row + 1 + i
        is_total = total_row and (i == len(rows) - 1)
        for pdf_idx, val in enumerate(row):
            excel_col = ncols - pdf_idx          # reversed mapping
            c = ws.cell(row=r, column=excel_col, value=val)
            c.alignment = center
            c.border = border
            if is_total:
                c.font = Font(name=FONT, bold=True, color=WHITE, size=font_size)
                c.fill = PatternFill(fill_type="solid", fgColor=brand)
            else:
                c.font = Font(name=FONT, size=font_size)
                if zebra and (i % 2 == 1):       # matches the PDF's even-row zebra
                    c.fill = PatternFill(fill_type="solid", fgColor=ZEBRA)
            if (pdf_idx in number_cols) and isinstance(val, (int, float)):
                c.number_format = "#,##0.00"
 
    if col_widths:
        for j, w in enumerate(reversed(col_widths), start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
    return start_row + 1 + len(rows)
 
 
def highlight_rows(ws, excel_rows, ncols, color):
    """Fill whole rows (e.g. the 'سداد' payment rows) with `color`, keeping borders."""
    color = _hex(color)
    for r in excel_rows:
        for j in range(1, ncols + 1):
            ws.cell(row=r, column=j).fill = PatternFill(fill_type="solid", fgColor=color)
 
 
def merge_on_total(ws, header_len, total_excel_row, pdf_from, pdf_to):
    """Merge a span on the totals row given PDF-order indices (inclusive)."""
    ncols = header_len
    c_a = ncols - pdf_to          # smaller excel col
    c_b = ncols - pdf_from
    ws.merge_cells(start_row=total_excel_row, start_column=c_a,
                   end_row=total_excel_row, end_column=c_b)
 
 
def write_totals(ws, brand, pairs, start_row):
    """Right-aligned (label | value) totals block; last row brand-coloured."""
    brand = _hex(brand)
    border = _border()
    for i, (label, value) in enumerate(pairs):
        r = start_row + i
        last = (i == len(pairs) - 1)
        lc = ws.cell(row=r, column=1, value=label)
        vc = ws.cell(row=r, column=2, value=float(value or 0))
        vc.number_format = "#,##0.00"
        for c in (lc, vc):
            c.border = border
            c.font = Font(name=FONT, bold=True, size=12, color=(WHITE if last else BLACK))
        lc.alignment = Alignment(horizontal="right", vertical="center")
        vc.alignment = Alignment(horizontal="center", vertical="center")
        lc.fill = PatternFill(fill_type="solid", fgColor=(brand if last else "EEF3FB"))
        if last:
            vc.fill = PatternFill(fill_type="solid", fgColor=brand)
    cur_a = ws.column_dimensions["A"].width or 0
    cur_b = ws.column_dimensions["B"].width or 0
    ws.column_dimensions["A"].width = max(cur_a, 24)
    ws.column_dimensions["B"].width = max(cur_b, 16)
 
 
def simple_report(buf, brand, title, date_from, date_to, header, rows,
                  number_cols=None, col_widths=None,
                  total_label=None, total_col=None, sheet="report"):
    """Title + table (+ optional single-column totals row). Covers most list reports."""
    wb, ws = new_sheet(sheet)
    ncols = len(header)
    r = write_title(ws, brand, title, date_from, date_to, ncols)
    data = [list(x) for x in rows]
    is_total = False
    if total_label is not None and total_col is not None:
        total_val = sum(float((x[total_col] or 0)) for x in rows) if rows else 0.0
        trow = [""] * ncols
        trow[0] = total_label
        trow[total_col] = total_val
        data.append(trow)
        is_total = True
    write_table(ws, brand, header, data, col_widths=col_widths, start_row=r,
                number_cols=number_cols, total_row=is_total)
    wb.save(buf)
 
 
# ---------------------------------------------------------------------------
# Shared response builder used by every export_* view (PDF or Excel).
# `chosen` is a list of tuples: (basename_without_ext, pdf_build_fn, xlsx_build_fn)
# ---------------------------------------------------------------------------
def deliver(request, chosen, fmt, zip_basename):
    from django.http import HttpResponse
    is_xlsx = (fmt == "xlsx")
    ext = "xlsx" if is_xlsx else "pdf"
    ct = XLSX_CT if is_xlsx else PDF_CT
 
    if len(chosen) == 1:
        base, pdf_b, xlsx_b = chosen[0]
        buf = io.BytesIO()
        (xlsx_b if is_xlsx else pdf_b)(buf)
        buf.seek(0)
        resp = HttpResponse(buf.getvalue(), content_type=ct)
        resp["Content-Disposition"] = f'attachment; filename="{base}.{ext}"'
        resp.set_cookie("export_download_done", "true", max_age=60)
        return resp
 
    zbuf = io.BytesIO()
    with zipfile.ZipFile(zbuf, "w", zipfile.ZIP_DEFLATED) as zf:
        for base, pdf_b, xlsx_b in chosen:
            b = io.BytesIO()
            (xlsx_b if is_xlsx else pdf_b)(b)
            zf.writestr(f"{base}.{ext}", b.getvalue())
    zbuf.seek(0)
    resp = HttpResponse(zbuf.getvalue(), content_type="application/zip")
    resp["Content-Disposition"] = f'attachment; filename="{zip_basename}.zip"'
    resp.set_cookie("export_download_done", "true", max_age=60)
    return resp
 