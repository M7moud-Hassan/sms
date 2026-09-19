"""
Overtime report -> styled Excel workbook (RTL, brand colours, live formulas).

Sheets
  1. الملخص    per-driver totals (formulas pull from the details sheet)
  2. التفاصيل  one row per driver/day that has overtime
  3. السندات حسب الأيام  every overtime day with the vouchers behind it

All labels live in LABELS so the language can be switched in one place.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    from .overtime import AMMAN_TZ
except ImportError:            # flat import when unit-testing
    from overtime import AMMAN_TZ

BRAND = "B45309"        # amber-700
BRAND_DARK = "7C2D12"
BAND = "FFF7ED"         # alternating driver groups
OFF_FILL = "E0F2FE"     # off-day cell highlight
TOTAL_FILL = "FDE68A"
SETTINGS_FILL = "F1F5F9"
GROUP_FILL = "FED7AA"   # day header rows on the vouchers sheet
FONT = "Arial"

LABELS = {
    'title': "تقرير العمل الإضافي للسائقين",
    'period': "الفترة من {} إلى {}",
    'drivers': "السائقون: {}",
    'summary_sheet': "الملخص",
    'details_sheet': "التفاصيل",
    'work_start': "بداية الدوام",
    'work_end': "نهاية الدوام",
    'before_mult': "معامل قبل الدوام",
    'after_mult': "معامل بعد الدوام",
    'off_mult': "معامل العطلة",
    'grace': "السماح (دقيقة)",
    'off_days': "أيام العطلة الأسبوعية: {}",
    'no_off_days': "لا يوجد",
    'not_counted': "غير محتسب",
    'work_day': "دوام",
    'off_day': "عطلة",
    'total': "الإجمالي",
    'no_data': "لا توجد ساعات عمل إضافي في هذه الفترة",
    'open_note': "ملاحظة: {} سند لم يتم تسليمه بعد ولم يدخل في الحساب",
    'summary_header': ["#", "السائق", "أيام العمل", "أيام العمل الإضافي", "أيام العطلة المعمولة",
                       "عدد السندات", "ساعات قبل الدوام", "ساعات بعد الدوام", "ساعات العطلة",
                       "إجمالي الساعات", "الساعات المحتسبة"],
    'details_header': ["#", "السائق", "التاريخ", "اليوم", "نوع اليوم", "بداية أول مهمة", "نهاية آخر مهمة",
                       "عدد السندات", "ساعات قبل الدوام", "ساعات بعد الدوام", "ساعات العطلة",
                       "إجمالي الساعات", "الساعات المحتسبة"],
    'vouchers_sheet': "السندات حسب الأيام",
    'vouchers_header': ["#", "رقم السند", "العميل", "نوع المركبة", "الماركة", "رقم اللوحة",
                        "موقع الاستلام", "وقت الاستلام", "موقع التسليم", "وقت التسليم",
                        "مدة السند (ساعة)", "ملاحظة"],
    'group': "{driver}   |   {weekday} {date}   |   {day_type}   |   من {first} إلى {last}   |   "
             "قبل الدوام: {before}   |   بعد الدوام: {after}   |   العطلة: {off}   |   المحتسبة: {weighted}",
    'role_begin': "سند مفتوح لأكثر من يوم: احتُسب وقت الاستلام فقط",
    'role_end': "سند مفتوح لأكثر من يوم: احتُسب وقت التسليم فقط",
    # indexed by python weekday(): Monday = 0
    'weekdays': ["الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت", "الأحد"],
}

THIN = Side(style="thin", color="D6D3D1")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
HOURS_FMT = '0.00;-0.00;"-"'
WEEK_ORDER = [6, 0, 1, 2, 3, 4, 5]   # Sunday first


def _fill(hex_color):
    return PatternFill("solid", start_color=hex_color, end_color=hex_color)


def _style_range(ws, row, first_col, last_col, *, fill=None, font=None, border=True, height=None):
    for c in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.alignment = CENTER
        if fill:
            cell.fill = fill
        if font:
            cell.font = font
        if border:
            cell.border = BORDER
    if height:
        ws.row_dimensions[row].height = height


def _prepare_sheet(ws, widths):
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True


def _banner(ws, row, ncols, text, fill, font, height):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    ws.cell(row=row, column=1, value=text)
    _style_range(ws, row, 1, ncols, fill=fill, font=font, border=False, height=height)


def _title_block(ws, ncols, date_from, date_to, driver_label):
    _banner(ws, 1, ncols, LABELS['title'], _fill(BRAND),
            Font(name=FONT, size=18, bold=True, color="FFFFFF"), 36)
    _banner(ws, 2, ncols,
            LABELS['period'].format(date_from.strftime('%Y-%m-%d'), date_to.strftime('%Y-%m-%d'))
            + "   |   " + LABELS['drivers'].format(driver_label),
            _fill(BRAND_DARK), Font(name=FONT, size=11, bold=True, color="FFFFFF"), 24)


def _header_row(ws, row, headers):
    for i, h in enumerate(headers, start=1):
        ws.cell(row=row, column=i, value=h)
    _style_range(ws, row, 1, len(headers), fill=_fill(BRAND),
                 font=Font(name=FONT, size=11, bold=True, color="FFFFFF"), height=34)


def build_overtime_xlsx(buf, report, cfg, date_from, date_to, driver_label):
    wb = Workbook()
    summary = wb.active
    summary.title = LABELS['summary_sheet']
    details = wb.create_sheet(LABELS['details_sheet'])
    vsheet = wb.create_sheet(LABELS['vouchers_sheet'])

    body_font = Font(name=FONT, size=11)
    bold_font = Font(name=FONT, size=11, bold=True)
    off_days = cfg.get('off_days') or set()
    off_txt = LABELS['off_days'].format(
        "، ".join(LABELS['weekdays'][d] for d in WEEK_ORDER if d in off_days) or LABELS['no_off_days'])
    before_lbl = LABELS['before_mult'] + ("" if cfg['count_before'] else f" ({LABELS['not_counted']})")
    after_lbl = LABELS['after_mult'] + ("" if cfg['count_after'] else f" ({LABELS['not_counted']})")

    header_row = 6
    first = header_row + 1

    # ================================================================ DETAILS
    d_head = LABELS['details_header']
    dn = len(d_head)
    _prepare_sheet(details, [6, 24, 13, 11, 10, 14, 14, 10, 13, 13, 13, 13, 14])
    _title_block(details, dn, date_from, date_to, driver_label)

    # settings strip (row 3): multiplier cells G3 / I3 / K3 are used by the formulas
    details.merge_cells("A3:B3")
    for ref, val in [("A3", LABELS['work_start']), ("C3", cfg['work_start']),
                     ("D3", LABELS['work_end']), ("E3", cfg['work_end']),
                     ("F3", before_lbl), ("G3", float(cfg['before_multiplier'])),
                     ("H3", after_lbl), ("I3", float(cfg['after_multiplier'])),
                     ("J3", LABELS['off_mult']), ("K3", float(cfg['off_day_multiplier'])),
                     ("L3", LABELS['grace']), ("M3", cfg['grace_minutes'])]:
        details[ref] = val
    _style_range(details, 3, 1, dn, fill=_fill(SETTINGS_FILL), font=bold_font, height=32)
    for ref in ("C3", "E3"):
        details[ref].number_format = "hh:mm AM/PM"
    for ref in ("G3", "I3", "K3"):
        details[ref].number_format = "0.00"
        details[ref].font = Font(name=FONT, size=11, bold=True, color="0000FF")
    _banner(details, 4, dn, off_txt, _fill(OFF_FILL), bold_font, 24)

    _header_row(details, header_row, d_head)
    r = first
    band_on, prev_driver = False, None
    for i, row in enumerate(report['rows'], start=1):
        if row['driver'] != prev_driver:
            band_on, prev_driver = not band_on, row['driver']
        values = [
            i, row['driver'], row['date'], LABELS['weekdays'][row['date'].weekday()],
            LABELS['off_day'] if row['is_off_day'] else LABELS['work_day'],
            row['first_start'].replace(tzinfo=None), row['last_end'].replace(tzinfo=None),
            row['vouchers'],
            float(row['before_hours']), float(row['after_hours']), float(row['off_day_hours']),
            f"=I{r}+J{r}+K{r}",
            f"=ROUND(I{r}*$G$3+J{r}*$I$3+K{r}*$K$3,2)",
        ]
        for c, v in enumerate(values, start=1):
            details.cell(row=r, column=c, value=v)
        _style_range(details, r, 1, dn, fill=_fill(BAND) if band_on else None, font=body_font, height=22)
        if row['is_off_day']:
            for c in (4, 5):
                details.cell(row=r, column=c).fill = _fill(OFF_FILL)
                details.cell(row=r, column=c).font = bold_font
        details.cell(row=r, column=3).number_format = "yyyy-mm-dd"
        details.cell(row=r, column=6).number_format = "hh:mm AM/PM"
        details.cell(row=r, column=7).number_format = "hh:mm AM/PM"
        for c in (9, 10, 11, 12, 13):
            details.cell(row=r, column=c).number_format = HOURS_FMT
        details.cell(row=r, column=13).font = bold_font
        r += 1
    last = r - 1

    if report['rows']:
        details.cell(row=r, column=1, value=LABELS['total'])
        details.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        details.cell(row=r, column=8, value=f"=SUM(H{first}:H{last})")
        for c in (9, 10, 11, 12, 13):
            col = get_column_letter(c)
            details.cell(row=r, column=c, value=f"=SUM({col}{first}:{col}{last})").number_format = HOURS_FMT
        _style_range(details, r, 1, dn, fill=_fill(TOTAL_FILL),
                     font=Font(name=FONT, size=12, bold=True, color=BRAND_DARK), height=26)
        details.auto_filter.ref = f"A{header_row}:{get_column_letter(dn)}{last}"
    else:
        _banner(details, r, dn, LABELS['no_data'], None, bold_font, 26)
    details.freeze_panes = f"A{first}"
    details.print_title_rows = f"{header_row}:{header_row}"

    # ================================================================ SUMMARY
    s_head = LABELS['summary_header']
    sn = len(s_head)
    _prepare_sheet(summary, [6, 26, 11, 13, 13, 11, 13, 13, 13, 13, 15])
    _title_block(summary, sn, date_from, date_to, driver_label)

    t = lambda v: v.strftime('%H:%M')
    _banner(summary, 3, sn,
            f"{LABELS['work_start']}: {t(cfg['work_start'])}   |   {LABELS['work_end']}: {t(cfg['work_end'])}"
            f"   |   {before_lbl}: ×{cfg['before_multiplier']:g}   |   {after_lbl}: ×{cfg['after_multiplier']:g}"
            f"   |   {LABELS['off_mult']}: ×{cfg['off_day_multiplier']:g}   |   {LABELS['grace']}: {cfg['grace_minutes']}",
            _fill(SETTINGS_FILL), bold_font, 28)
    _banner(summary, 4, sn, off_txt, _fill(OFF_FILL), bold_font, 24)

    _header_row(summary, header_row, s_head)
    ds = f"'{LABELS['details_sheet']}'!"
    has_rows = bool(report['rows'])
    rng = lambda col: f"{ds}${col}${first}:${col}${last}"
    r = first
    for i, drv in enumerate(report['drivers'], start=1):
        n = f"$B{r}"
        if has_rows:
            ot_days = f"=COUNTIF({rng('B')},{n})"
            off_worked = f"=COUNTIFS({rng('B')},{n},{rng('E')},\"{LABELS['off_day']}\")"
            before = f"=SUMIF({rng('B')},{n},{rng('I')})"
            after = f"=SUMIF({rng('B')},{n},{rng('J')})"
            off_h = f"=SUMIF({rng('B')},{n},{rng('K')})"
            weighted = f"=SUMIF({rng('B')},{n},{rng('M')})"
        else:
            ot_days = off_worked = before = after = off_h = weighted = 0
        values = [i, drv['driver'], drv['work_days'], ot_days, off_worked, drv['vouchers'],
                  before, after, off_h, f"=G{r}+H{r}+I{r}", weighted]
        for c, v in enumerate(values, start=1):
            summary.cell(row=r, column=c, value=v)
        _style_range(summary, r, 1, sn, fill=_fill(BAND) if i % 2 == 0 else None, font=body_font, height=24)
        for c in (7, 8, 9, 10, 11):
            summary.cell(row=r, column=c).number_format = HOURS_FMT
        summary.cell(row=r, column=11).font = bold_font
        r += 1
    s_last = r - 1

    if report['drivers']:
        summary.cell(row=r, column=1, value=LABELS['total'])
        summary.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        for c in range(3, sn + 1):
            col = get_column_letter(c)
            cell = summary.cell(row=r, column=c, value=f"=SUM({col}{first}:{col}{s_last})")
            if c >= 7:
                cell.number_format = HOURS_FMT
        _style_range(summary, r, 1, sn, fill=_fill(TOTAL_FILL),
                     font=Font(name=FONT, size=12, bold=True, color=BRAND_DARK), height=28)
    else:
        _banner(summary, r, sn, LABELS['no_data'], None, bold_font, 26)

    if report.get('open_vouchers'):
        _banner(summary, r + 2, sn, LABELS['open_note'].format(report['open_vouchers']), None,
                Font(name=FONT, size=10, italic=True, color="B91C1C"), None)

    summary.freeze_panes = f"A{first}"
    summary.print_title_rows = f"{header_row}:{header_row}"

    _build_vouchers_sheet(vsheet, report, date_from, date_to, driver_label, off_txt, header_row)
    wb.save(buf)


def _local_naive(dt):
    return dt.astimezone(AMMAN_TZ).replace(tzinfo=None) if dt else None


def _build_vouchers_sheet(ws, report, date_from, date_to, driver_label, off_txt, header_row):
    head = LABELS['vouchers_header']
    n = len(head)
    body_font = Font(name=FONT, size=11)
    bold_font = Font(name=FONT, size=11, bold=True)
    _prepare_sheet(ws, [6, 11, 24, 14, 13, 14, 26, 23, 26, 23, 13, 44])
    _title_block(ws, n, date_from, date_to, driver_label)
    _banner(ws, 3, n, off_txt, _fill(OFF_FILL), bold_font, 24)
    _header_row(ws, header_row, head)

    details = report.get('voucher_details') or {}
    r = header_row + 1
    if not report['rows']:
        _banner(ws, r, n, LABELS['no_data'], None, bold_font, 26)

    for row in report['rows']:
        first, last = _local_naive(row['first_start']), _local_naive(row['last_end'])
        last_txt = last.strftime('%H:%M') if last.date() == row['date'] else last.strftime('%H:%M  %Y-%m-%d')
        _banner(ws, r, n, LABELS['group'].format(
            driver=row['driver'], weekday=LABELS['weekdays'][row['date'].weekday()],
            date=row['date'].strftime('%Y-%m-%d'),
            day_type=LABELS['off_day'] if row['is_off_day'] else LABELS['work_day'],
            first=first.strftime('%H:%M'), last=last_txt,
            before=row['before_hours'], after=row['after_hours'], off=row['off_day_hours'],
            weighted=row['weighted_hours']),
            _fill(GROUP_FILL), Font(name=FONT, size=11, bold=True, color=BRAND_DARK), 26)
        for c in range(1, n + 1):
            ws.cell(row=r, column=c).border = BORDER
        r += 1

        items = sorted(row['voucher_items'],
                       key=lambda it: (details.get(it[0], {}).get('receipt_from_time') is None,
                                       _local_naive(details.get(it[0], {}).get('receipt_from_time')) or first))
        for i, (vid, role) in enumerate(items, start=1):
            v = details.get(vid, {})
            car = lambda k: v.get(k) or ""
            values = [i, v.get('number') or vid, car('car_owner__name'), car('car_type'), car('car_mark'),
                      car('car_num'), car('receipt_from_location'), _local_naive(v.get('receipt_from_time')),
                      car('receipt_to_location'), _local_naive(v.get('receipt_to_time')),
                      f"=ROUND((J{r}-H{r})*24,2)",
                      LABELS['role_begin'] if role == 'begin' else LABELS['role_end'] if role == 'end' else ""]
            for c, val in enumerate(values, start=1):
                ws.cell(row=r, column=c, value=val)
            _style_range(ws, r, 1, n, fill=_fill(BAND) if i % 2 == 0 else None, font=body_font, height=22)
            for c in (8, 10):
                ws.cell(row=r, column=c).number_format = "yyyy-mm-dd hh:mm AM/PM"
            ws.cell(row=r, column=11).number_format = "0.00"
            if role != 'trip':
                note = ws.cell(row=r, column=12)
                note.font = Font(name=FONT, size=10, bold=True, color="B91C1C")
                for c in ((8, 12) if role == 'begin' else (10, 12)):
                    ws.cell(row=r, column=c).fill = _fill("FEE2E2")
            r += 1
        r += 1   # blank line between days

    ws.freeze_panes = f"A{header_row + 1}"
    ws.print_title_rows = f"{header_row}:{header_row}"
